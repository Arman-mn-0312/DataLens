import sys
import os
import io
import json
import uuid
import pandas as pd
import openpyxl
from pathlib import Path

root_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(root_dir))

from app import app
from auth.security import create_access_token
from services.dataset_manager import DatasetManager
from services.report_cache_service import ReportCacheService
from services.report_export_service import ReportExportService
from services.upload_storage import get_user_upload_path
from services.datatype_service import TYPE_CHECKERS, classify_column


def validate_datatype_compatibility():
    samples = [
        pd.Series(["", "NaN", "null", "0", "$1,203.50", "20%", "TRUE", "yes", "n", "2024-01-02", "13/01/2024", "not date", None], dtype=object),
        pd.Series([True, False, True], dtype=bool),
        pd.Series([1.0, None, -2.0]),
        pd.Series(pd.to_datetime(["2024-01-01", None, "2024-02-01"])),
    ]
    for series in samples:
        expected = []
        for value in series.dropna():
            detected = None
            for type_name, checker in TYPE_CHECKERS.items():
                if checker(value):
                    detected = type_name
                    break
            expected.append(detected)
        actual, _ = classify_column(series)
        assert actual == expected, f"Datatype classification changed for {series.dtype}: {actual} != {expected}"

def run_comprehensive_validation():
    print("==================================================")
    validate_datatype_compatibility()
    print("Datatype classifier matches prior behavior for mixed and native dtypes.")
    print("DATALENS COMPREHENSIVE END-TO-END VALIDATION")
    print("==================================================")

    client = app.test_client()
    auth_headers = {"Authorization": f"Bearer {create_access_token('qa-comprehensive-validation')}"}

    # Step 1: Upload CSV
    sample_csv_path = os.path.join(root_dir, "data", "sample", "combined_dataset.csv")
    with open(sample_csv_path, "rb") as f:
        upload_resp = client.post(
            "/upload",
            data={"file": (f, "combined_dataset.csv")},
            content_type="multipart/form-data",
            headers=auth_headers,
        )
    print(f"1. Upload Response: {upload_resp.status_code}")
    assert upload_resp.status_code == 200, f"Upload failed: {upload_resp.data}"
    upload_data = upload_resp.get_json()
    assert upload_data["success"] is True
    print(f"   Uploaded filename: {upload_data['filename']}")

    # Step 2: Fetch all reports to simulate user viewing pages
    print("2. Simulating User Workflow (Loading all Report Pages)...")
    for rep in ["overview", "missing", "duplicate", "datatype", "outlier", "dashboard"]:
        r = client.get(f"/reports/{rep}?filename=combined_dataset.csv", headers=auth_headers)
        assert r.status_code == 200, f"Failed to fetch {rep}: {r.status_code}"
        payload = r.get_json()
        print(f"   [OK] Fetched /reports/{rep} successfully (Cached: {ReportCacheService.has_report(rep, 'combined_dataset.csv', session_key='qa-comprehensive-validation')})")

    # Verify cache values
    overview_data = ReportCacheService.get_report("overview", "combined_dataset.csv", session_key="qa-comprehensive-validation")
    dashboard_data = ReportCacheService.get_report("dashboard", "combined_dataset.csv", session_key="qa-comprehensive-validation")
    assert overview_data is not None, "Expected overview report to be cached"
    assert dashboard_data is not None, "Expected dashboard report to be cached"

    summary = dashboard_data.get("summary") or {}
    expected_score = summary.get("health_score")
    expected_status = summary.get("status")
    print(f"   Extracted Analysis Score: {expected_score}/100, Status: {expected_status}")

    # A malformed replacement with the same name must leave the active report usable.
    invalid_replacement = client.post(
        "/upload",
        data={"file": (io.BytesIO(b""), "combined_dataset.csv")},
        content_type="multipart/form-data",
        headers=auth_headers,
    )
    assert invalid_replacement.status_code == 400
    preserved = client.get("/reports/overview?filename=combined_dataset.csv", headers=auth_headers)
    assert preserved.status_code == 200
    assert preserved.get_json()["total_rows"] == 12
    print("   Invalid replacement rejected; the previous dataset remains available.")

    # Step 3: Test PDF Generation & Verification
    print("3. Testing PDF Generation Endpoint (POST /reports/generate/pdf)...")
    pdf_resp = client.post("/reports/generate/pdf?filename=combined_dataset.csv", headers=auth_headers)
    assert pdf_resp.status_code == 200
    assert "application/pdf" in pdf_resp.content_type
    assert pdf_resp.data.startswith(b"%PDF")
    content_disp = pdf_resp.headers.get("Content-Disposition", "")
    assert "DataLens_combined_dataset_report.pdf" in content_disp
    print(f"   [OK] PDF Exported: Size={len(pdf_resp.data)} bytes, Disposition={content_disp}")

    # Step 4: Test Excel Generation & Sheet Verification
    print("4. Testing Excel Generation Endpoint (POST /reports/generate/excel)...")
    excel_resp = client.post("/reports/generate/excel?filename=combined_dataset.csv", headers=auth_headers)
    assert excel_resp.status_code == 200
    assert "spreadsheetml" in excel_resp.content_type
    assert excel_resp.data.startswith(b"PK")
    content_disp_xl = excel_resp.headers.get("Content-Disposition", "")
    assert "DataLens_combined_dataset_report.xlsx" in content_disp_xl
    print(f"   [OK] Excel Exported: Size={len(excel_resp.data)} bytes, Disposition={content_disp_xl}")

    # Inspect Excel Worksheets
    wb = openpyxl.load_workbook(io.BytesIO(excel_resp.data))
    expected_sheets = [
        "Executive Summary",
        "Dataset Overview",
        "Missing Values",
        "Duplicate Records",
        "Datatype Validation",
        "Outlier Detection",
        "Recommendations"
    ]
    print(f"   Excel Sheet Names: {wb.sheetnames}")
    for sheet_name in expected_sheets:
        assert sheet_name in wb.sheetnames, f"Missing expected sheet: {sheet_name}"
        ws = wb[sheet_name]
        assert ws.max_row >= 4, f"Sheet {sheet_name} has insufficient rows ({ws.max_row})"
        print(f"     - Sheet '{sheet_name}': {ws.max_row} rows, {ws.max_column} cols [OK]")

    # Verify values inside Excel Sheet 1 (Executive Summary)
    ws_exec = wb["Executive Summary"]
    found_score = False
    for r in range(1, 12):
        cell_val = str(ws_exec[f"B{r}"].value or "")
        if f"{expected_score} / 100" in cell_val:
            found_score = True
            print(f"     - Verified Health Score in Sheet 1: {cell_val} matches single source of truth!")
            break
    assert found_score, f"Health Score {expected_score} / 100 not found in Sheet 1"

    # Step 5: Test Export without pre-fetching (Cold cache on fresh upload)
    print("5. Testing Cold Export (Generating PDF/Excel without prior page loads)...")
    ReportCacheService.clear_cache(session_key="qa-comprehensive-validation")
    cold_pdf_resp = client.post("/reports/generate/pdf?filename=combined_dataset.csv", headers=auth_headers)
    assert cold_pdf_resp.status_code == 200
    assert cold_pdf_resp.data.startswith(b"%PDF")
    print(f"   [OK] Cold PDF Export succeeded: Size={len(cold_pdf_resp.data)} bytes")

    # Step 6: Test Error handling for invalid dataset
    print("6. Testing Error Handling...")
    err_resp = client.post("/reports/generate/pdf?filename=non_existent_123.csv", headers=auth_headers)
    assert err_resp.status_code == 400, f"Expected 400, got {err_resp.status_code}"
    print(f"   [OK] Invalid file correctly rejected with status {err_resp.status_code}")

    # Step 7: Verify that replacing the active file invalidates all report state
    # and that every report is generated against the newly uploaded CSV.
    replacement_filename = f"qa_reupload_{uuid.uuid4().hex}.csv"
    try:
        replacement = client.post(
            "/upload",
            data={"file": (io.BytesIO(b"item,score\nA,1\nB,100\nB,100\n"), replacement_filename)},
            content_type="multipart/form-data",
            headers=auth_headers,
        )
        assert replacement.status_code == 200, replacement.get_json()
        replacement_overview = None
        for rep in ["overview", "missing", "duplicate", "datatype", "outlier", "dashboard"]:
            response = client.get(
                f"/reports/{rep}?filename={replacement_filename}",
                headers=auth_headers,
            )
            assert response.status_code == 200, f"Reupload report {rep} failed: {response.get_data(as_text=True)}"
            if rep == "overview":
                replacement_overview = response.get_json()
        assert replacement_overview["dataset_name"] == replacement_filename
        assert replacement_overview["total_rows"] == 3

        same_name_replacement = client.post(
            "/upload",
            data={"file": (io.BytesIO(b"item,score\nC,7\nD,8\n"), replacement_filename)},
            content_type="multipart/form-data",
            headers=auth_headers,
        )
        assert same_name_replacement.status_code == 200, same_name_replacement.get_json()
        for rep in ["overview", "missing", "duplicate", "datatype", "outlier", "dashboard"]:
            response = client.get(
                f"/reports/{rep}?filename={replacement_filename}",
                headers=auth_headers,
            )
            assert response.status_code == 200, f"Same-name report {rep} failed: {response.get_data(as_text=True)}"
            if rep == "overview":
                assert response.get_json()["total_rows"] == 2
        print("7. Reupload check passed for new and same-name CSVs; every report used the latest data.")
        deleted = client.delete(
            f"/upload?filename={replacement_filename}",
            headers=auth_headers,
        )
        assert deleted.status_code == 200, deleted.get_json()
        assert not os.path.exists(get_user_upload_path(replacement_filename, "qa-comprehensive-validation"))
        print("8. Explicit delete removes the server-side upload.")

        # Excel and JSON imports should enter the exact same analysis flow.
        tabular_rows = [{"item": "A", "score": 1}, {"item": "B", "score": 100}]
        workbook = io.BytesIO()
        pd.DataFrame(tabular_rows).to_excel(workbook, index=False)
        format_payloads = [
            (f"qa_format_{uuid.uuid4().hex}.xlsx", workbook.getvalue()),
            (f"qa_format_{uuid.uuid4().hex}.json", json.dumps(tabular_rows).encode("utf-8")),
        ]
        try:
            for format_name, format_bytes in format_payloads:
                imported = client.post(
                    "/upload",
                    data={"file": (io.BytesIO(format_bytes), format_name)},
                    content_type="multipart/form-data",
                    headers=auth_headers,
                )
                assert imported.status_code == 200, imported.get_json()
                for report_name in ["overview", "missing", "duplicate", "datatype", "outlier", "dashboard"]:
                    report_response = client.get(
                        f"/reports/{report_name}?filename={format_name}",
                        headers=auth_headers,
                    )
                    assert report_response.status_code == 200, report_response.get_data(as_text=True)
                    if report_name == "overview":
                        assert report_response.get_json()["total_rows"] == 2
                print(f"9. {format_name.rsplit('.', 1)[-1].upper()} upload and all report operations passed.")
        finally:
            for format_name, _ in format_payloads:
                format_path = get_user_upload_path(format_name, "qa-comprehensive-validation")
                if os.path.exists(format_path):
                    os.remove(format_path)
    finally:
        DatasetManager.clear_cache(session_key="qa-comprehensive-validation")
        ReportCacheService.clear_cache(session_key="qa-comprehensive-validation")
        for uploaded_name in ("combined_dataset.csv", replacement_filename):
            uploaded_path = get_user_upload_path(uploaded_name, "qa-comprehensive-validation")
            if uploaded_path and os.path.exists(uploaded_path):
                os.remove(uploaded_path)

    print("==================================================")
    print("ALL VERIFICATIONS COMPLETED WITH 100% SUCCESS!")
    print("==================================================")
    return True

if __name__ == "__main__":
    success = run_comprehensive_validation()
    sys.exit(0 if success else 1)
